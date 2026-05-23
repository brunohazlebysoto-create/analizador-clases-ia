import os
import time
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def format_timestamp(seconds):
    """Convierte segundos a formato hh:mm:ss"""
    h = int(seconds // 3600)
    m = int((seconds % 3600) // 60)
    s = int(seconds % 60)
    return f"{h:02d}:{m:02d}:{s:02d}"

def set_cell_background(cell, hex_color):
    """Establece el color de fondo de una celda de tabla en Word"""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), hex_color)
    shd.set(qn('w:val'), 'clear')
    tcPr.append(shd)

def set_paragraph_callout_style(paragraph, color_hex="1A365D", bg_hex="F0F4F8"):
    """Dibuja un borde izquierdo grueso (Callout) y fondo sombreado en un párrafo de Word"""
    pPr = paragraph._element.get_or_add_pPr()
    
    # Borde izquierdo (pBdr)
    pBdr = OxmlElement('w:pBdr')
    left = OxmlElement('w:left')
    left.set(qn('w:val'), 'single')
    left.set(qn('w:sz'), '24') # 24 octavos de punto = 3pt de grosor
    left.set(qn('w:space'), '12') # Margen interno (padding)
    left.set(qn('w:color'), color_hex)
    pBdr.append(left)
    pPr.append(pBdr)
    
    # Sombreado de fondo (shd)
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), bg_hex)
    shd.set(qn('w:val'), 'clear')
    pPr.append(shd)

def _process_inline_markdown(paragraph, text, text_color, default_italic=False, font_size=11):
    """
    Parsea recursivamente/secuencialmente texto en formato Markdown simple
    con soporte para **negrita** e *itálica*.
    """
    bold_parts = text.split("**")
    for b_idx, b_part in enumerate(bold_parts):
        is_bold = (b_idx % 2 == 1)
        
        italic_parts = b_part.split("*")
        for i_idx, i_part in enumerate(italic_parts):
            is_italic = (i_idx % 2 == 1) or default_italic
            
            if not i_part:
                continue
                
            run = paragraph.add_run(i_part)
            run.font.name = 'Arial'
            run.font.size = Pt(font_size)
            run.font.color.rgb = text_color
            
            if is_bold:
                run.font.bold = True
            if is_italic:
                run.font.italic = True

def _add_markdown_table_to_docx(doc, table_lines, primary_color, text_color):
    """Parsea una tabla en formato Markdown y la añade como tabla nativa de Word"""
    rows_data = []
    for line in table_lines:
        cells = [c.strip() for c in line.split('|')]
        # Quitar los extremos vacíos si empieza/termina con '|'
        if line.startswith('|'):
            cells = cells[1:]
        if line.endswith('|'):
            cells = cells[:-1]
            
        # Omitir fila separadora (ej. |---|---|)
        if cells and all(all(ch in '-: ' for ch in cell) for cell in cells):
            continue
            
        rows_data.append(cells)
        
    if not rows_data:
        return
        
    num_cols = max(len(r) for r in rows_data)
    table = doc.add_table(rows=0, cols=num_cols)
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER
    table.autofit = True
    table.style = 'Table Grid'
    
    for r_idx, row in enumerate(rows_data):
        row_cells = table.add_row().cells
        for c_idx, val in enumerate(row):
            if c_idx < len(row_cells):
                cell = row_cells[c_idx]
                
                # Margen de párrafo interno
                p = cell.paragraphs[0]
                p.text = ""
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(4)
                
                is_header = (r_idx == 0)
                cell_color = RGBColor(255, 255, 255) if is_header else text_color
                _process_inline_markdown(p, val, cell_color, font_size=10)
                
                if is_header:
                    for run in p.runs:
                        run.font.bold = True
                        
                # Aplicar sombreado
                if is_header:
                    set_cell_background(cell, "1A365D") # Azul Marino cabecera
                elif r_idx % 2 == 1:
                    set_cell_background(cell, "F7FAFC") # Gris alternante zebra
                    
    # Espaciado después de la tabla
    p_space = doc.add_paragraph()
    p_space.paragraph_format.space_after = Pt(10)

def create_word_report(slides_data, class_title="Reporte de Clase Automatizado", output_filename="reporte_clase.docx"):
    """
    Genera un documento Word (.docx) formal y bien estructurado.
    Inserta la imagen de la diapositiva y, debajo, la explicación sintetizada.
    """
    doc = Document()
    
    # Configuración de márgenes estándar (1 pulgada = 2.54 cm)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        
    # Paleta de colores ejecutiva (Azul Marino y Gris Oxford)
    primary_color = RGBColor(26, 54, 93)      # #1A365D - Encabezados principales
    secondary_color = RGBColor(74, 85, 104)   # #4A5568 - Subtítulos
    text_color = RGBColor(45, 55, 72)         # #2D3748 - Texto normal
    
    # === PORTADA ELEGANTE ===
    for _ in range(3):
        doc.add_paragraph()
        
    # Título Principal (Grande, Bold)
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run(class_title)
    title_run.font.name = 'Arial'
    title_run.font.size = Pt(28)
    title_run.font.bold = True
    title_run.font.color.rgb = primary_color
    title_p.paragraph_format.space_after = Pt(12)
    
    # Subtítulo
    subtitle_p = doc.add_paragraph()
    subtitle_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_run = subtitle_p.add_run("Reporte de Síntesis y Verificación Científica de Clase")
    subtitle_run.font.name = 'Arial'
    subtitle_run.font.size = Pt(14)
    subtitle_run.font.italic = True
    subtitle_run.font.color.rgb = secondary_color
    subtitle_p.paragraph_format.space_after = Pt(36)
    
    # Línea horizontal divisoria gruesa
    p_line = doc.add_paragraph()
    p_line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_line.paragraph_format.space_after = Pt(48)
    line_run = p_line.add_run("―" * 40)
    line_run.font.bold = True
    line_run.font.color.rgb = primary_color
    
    # Espaciado
    for _ in range(4):
        doc.add_paragraph()
        
    # Metadatos del Documento
    meta_p = doc.add_paragraph()
    meta_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta_p.paragraph_format.space_after = Pt(6)
    meta_label = meta_p.add_run("Creado por: ")
    meta_label.font.bold = True
    meta_label.font.size = Pt(10)
    meta_label.font.color.rgb = secondary_color
    
    meta_val = meta_p.add_run("Analizador de Clases Inteligente (Gemini AI)")
    meta_val.font.size = Pt(10)
    meta_val.font.color.rgb = text_color
    
    date_p = doc.add_paragraph()
    date_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    date_p.paragraph_format.space_after = Pt(6)
    date_label = date_p.add_run("Fecha de Generación: ")
    date_label.font.bold = True
    date_label.font.size = Pt(10)
    date_label.font.color.rgb = secondary_color
    
    date_val = date_p.add_run(time.strftime('%d/%m/%Y a las %H:%M'))
    date_val.font.size = Pt(10)
    date_val.font.color.rgb = text_color
    
    # Salto de página para comenzar el reporte en la página 2
    doc.add_page_break()
    
    # Agregar contenido de las diapositivas
    for idx, slide in enumerate(slides_data):
        # Título de la diapositiva: "Diapositiva 1 (00:00:00 - 00:02:15)"
        h_p = doc.add_paragraph()
        h_p.paragraph_format.space_before = Pt(18)
        h_p.paragraph_format.space_after = Pt(12)
        h_p.paragraph_format.keep_with_next = True
        
        start_str = format_timestamp(slide["start_time"])
        end_str = format_timestamp(slide["end_time"])
        
        h_run = h_p.add_run(f"Diapositiva {slide['id']} ({start_str} - {end_str})")
        h_run.font.name = 'Arial'
        h_run.font.size = Pt(14)
        h_run.font.bold = True
        h_run.font.color.rgb = primary_color
        
        # Imagen de la Diapositiva (Pantallazo)
        image_path = slide["image_path"]
        if os.path.exists(image_path):
            img_p = doc.add_paragraph()
            img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            img_p.paragraph_format.space_after = Pt(14)
            try:
                img_p.add_run().add_picture(image_path, width=Inches(5.5))
            except Exception as e:
                err_run = img_p.add_run(f"[Error al insertar imagen {slide['image_name']}: {str(e)}]")
                err_run.font.italic = True
                err_run.font.color.rgb = RGBColor(150, 0, 0)
        else:
            no_img_p = doc.add_paragraph()
            no_img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            no_run = no_img_p.add_run(f"[Imagen no encontrada: {slide['image_name']}]")
            no_run.font.italic = True
            no_run.font.color.rgb = secondary_color
            
        # Título para la explicación
        exp_title_p = doc.add_paragraph()
        exp_title_p.paragraph_format.space_before = Pt(6)
        exp_title_p.paragraph_format.space_after = Pt(4)
        exp_title_p.paragraph_format.keep_with_next = True
        exp_title_run = exp_title_p.add_run("Explicación Integrada:")
        exp_title_run.font.name = 'Arial'
        exp_title_run.font.size = Pt(11)
        exp_title_run.font.bold = True
        exp_title_run.font.color.rgb = secondary_color
        
        # Contenido de la Explicación (procesando marcas markdown de Gemini)
        explanation = slide.get("explanation", "No se proporcionó explicación para esta diapositiva.")
        lines = explanation.split('\n')
        
        i = 0
        while i < len(lines):
            line = lines[i].strip()
            if not line:
                i += 1
                continue
            
            # 1. Detectar tablas en Markdown
            if line.startswith('|'):
                table_lines = []
                while i < len(lines) and lines[i].strip().startswith('|'):
                    table_lines.append(lines[i].strip())
                    i += 1
                _add_markdown_table_to_docx(doc, table_lines, primary_color, text_color)
                continue
                
            # 2. Detectar Callouts (Nota de Atención o Alerta Clínica)
            is_callout = False
            callout_type = None
            
            p_text_lower = line.lower()
            if "nota de atención:" in p_text_lower or "nota de atencion:" in p_text_lower:
                is_callout = True
                callout_type = "atencion"
            elif "alerta clínica:" in p_text_lower or "alerta clinica:" in p_text_lower:
                is_callout = True
                callout_type = "alerta"
                
            if is_callout:
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(8)
                p.paragraph_format.space_after = Pt(8)
                p.paragraph_format.left_indent = Inches(0.2)
                p.paragraph_format.line_spacing = 1.15
                
                if callout_type == "alerta":
                    set_paragraph_callout_style(p, color_hex="9B2C2C", bg_hex="FFF5F5")
                else:
                    set_paragraph_callout_style(p, color_hex="1A365D", bg_hex="F0F4F8")
                
                # Limpiar los asteriscos externos del Callout
                clean_text = line
                if clean_text.startswith("*"):
                    clean_text = clean_text[1:]
                if clean_text.endswith("*"):
                    clean_text = clean_text[:-1]
                clean_text = clean_text.strip()
                
                # Identificar el prefijo y ponerlo en negrita
                prefix_candidate = ""
                if callout_type == "alerta":
                    colon_idx = clean_text.lower().find("alerta cl\xednica:")
                    if colon_idx == -1:
                        colon_idx = clean_text.lower().find("alerta clinica:")
                    if colon_idx != -1:
                        end_prefix = colon_idx + len("alerta clínica:")
                        prefix_candidate = clean_text[:end_prefix]
                else:
                    colon_idx = clean_text.lower().find("nota de atenci\xf3n:")
                    if colon_idx == -1:
                        colon_idx = clean_text.lower().find("nota de atencion:")
                    if colon_idx != -1:
                        end_prefix = colon_idx + len("nota de atención:")
                        prefix_candidate = clean_text[:end_prefix]
                
                if prefix_candidate:
                    body_text = clean_text[len(prefix_candidate):].strip()
                    run_pref = p.add_run(prefix_candidate + " ")
                    run_pref.font.name = 'Arial'
                    run_pref.font.size = Pt(10.5)
                    run_pref.font.bold = True
                    run_pref.font.italic = True
                    if callout_type == "alerta":
                        run_pref.font.color.rgb = RGBColor(155, 44, 44)
                    else:
                        run_pref.font.color.rgb = RGBColor(26, 54, 93)
                    
                    _process_inline_markdown(p, body_text, text_color, default_italic=True)
                else:
                    _process_inline_markdown(p, clean_text, text_color, default_italic=True)
                
                i += 1
                continue
                
            # 3. Detectar encabezados
            if line.startswith("### "):
                header_title = line.replace("### ", "").strip()
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(14)
                p.paragraph_format.space_after = Pt(6)
                p.paragraph_format.keep_with_next = True
                
                if any(term in header_title for term in ["Corroboración Científica", "Corroboracion Cientifica", "Aclaración de Errores", "Aclaracion de Errores"]):
                    run = p.add_run("🔬 " + header_title)
                else:
                    run = p.add_run("▪ " + header_title)
                    
                run.font.name = 'Arial'
                run.font.size = Pt(12.5)
                run.font.bold = True
                run.font.color.rgb = primary_color
                
            elif line.startswith("## "):
                header_title = line.replace("## ", "").strip()
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(16)
                p.paragraph_format.space_after = Pt(8)
                p.paragraph_format.keep_with_next = True
                
                run = p.add_run("◈ " + header_title)
                run.font.name = 'Arial'
                run.font.size = Pt(14)
                run.font.bold = True
                run.font.color.rgb = primary_color
                
            elif line.startswith("# "):
                header_title = line.replace("# ", "").strip()
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(18)
                p.paragraph_format.space_after = Pt(10)
                p.paragraph_format.keep_with_next = True
                
                run = p.add_run("❖ " + header_title)
                run.font.name = 'Arial'
                run.font.size = Pt(16)
                run.font.bold = True
                run.font.color.rgb = primary_color
                
            # 4. Detectar viñetas de lista
            elif line.startswith("- ") or line.startswith("* "):
                p = doc.add_paragraph(style='List Bullet')
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(3)
                p.paragraph_format.line_spacing = 1.15
                bullet_text = line[2:]
                _process_inline_markdown(p, bullet_text, text_color)
                
            # 5. Detectar listas numeradas
            elif line and line[0].isdigit() and len(line) > 2 and line[1:3] == ". ":
                p = doc.add_paragraph(style='List Number')
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(3)
                p.paragraph_format.line_spacing = 1.15
                num_text = line[3:]
                _process_inline_markdown(p, num_text, text_color)
                
            # 6. Párrafo normal
            else:
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(6)
                p.paragraph_format.line_spacing = 1.15
                _process_inline_markdown(p, line, text_color)
                
            i += 1
            
        # Salto de página para la siguiente diapositiva
        if idx < len(slides_data) - 1:
            doc.add_page_break()
            
    # Guardar el documento
    doc.save(output_filename)
    return output_filename
